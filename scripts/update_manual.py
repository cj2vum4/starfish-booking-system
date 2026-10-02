"""Update the original manual in place, preserving existing run formatting."""
from pathlib import Path
import json, shutil
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
MANUAL = ROOT.parent / 'Starfish_LINE_OA_預約系統建置手冊_v1.0.docx'
BACKUP = ROOT / 'docs/backups' / MANUAL.name
if not BACKUP.exists():
    BACKUP.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(MANUAL, BACKUP)
state = json.loads((ROOT/'docs/stages.json').read_text(encoding='utf-8'))
doc = Document(MANUAL)

def replace_text(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return
    if old in p.text:
        text = p.text.replace(old, new)
        if p.runs:
            p.runs[0].text = text
            for r in p.runs[1:]: r.text = ''
        else: p.add_run(text)

for table in doc.tables:
    if len(table.columns)==4 and table.cell(0,0).text=='階段':
        for row in table.rows[1:]:
            key = row.cells[0].text.strip()
            if key in state['stages']:
                p=row.cells[3].paragraphs[0]
                replace_text(p,p.text,state['stages'][key]['status'])

completed = [
    '建立 / 確認 LINE Official Account。',
    '在 LINE Developers 建立 Provider。',
    '建立 Messaging API Channel。',
    '建立 .env.example，只放變數名稱，不提交 secret。',
    '確認 LINE OA 管理權限。',
    '登入 LINE Developers，確認 / 建立 Provider。',
    '建立 users table。',
    '完成 signature 驗證。',
    '用自己的 LINE 帳號加入 / 解除 / 重加，確認 users 正確 upsert。',
    'users.line_user_id UNIQUE。',
    '所有時間使用 timestamptz 儲存，前端以 Asia/Taipei 顯示。'
]
section = ''
for p in doc.paragraphs:
    if p.text.startswith('3. P1'): section='P1'
    elif p.text.startswith('4. P2'): section='P2'
    if p.text.startswith('☐') and (section=='P1' or any(x in p.text for x in completed)):
        # Front-end Taipei display is not implemented yet.
        if '前端以 Asia/Taipei' not in p.text: replace_text(p,'☐','☒')

heading='22. 實際施工與驗收進度'
found=False
for p in list(doc.paragraphs):
    if p.text==heading: found=True
    if found:
        p._element.getparent().remove(p._element)
# Appendix uses paragraphs only, so previous appendix removal is deterministic.
doc.add_heading(heading,level=1)
doc.add_paragraph('更新日期：'+state['updated']+'。本節依程式、雲端設定與實際測試更新。已完成僅代表該階段驗收通過，整體系統尚未正式上線。')
for key,value in state['stages'].items():
    p=doc.add_paragraph()
    p.add_run(f'{key}　{value["status"]}　').bold=True
    p.add_run(value['evidence'])
doc.add_paragraph('程式與詳細證據：同目錄 starfish-booking-system。階段狀態保存在 docs/stages.json；測試證據位於 docs/evidence，部署紀錄位於 docs/PROGRESS.md。')
doc.save(MANUAL)
print('Updated:', MANUAL)
